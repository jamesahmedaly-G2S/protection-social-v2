# -*- coding: utf-8 -*-
"""Comparaison des populations de salariés entre PAIE et DSN.

Trois niveaux de comparaison, tous construits sur la même fonction générique
`comparer()` appliquée à deux populations déjà chargées {clé société: {matricule: (nom, prénom)}} :
  - reconstruits       : à partir des classeurs déjà reconstruits (output/xlsx/) ;
  - sources brutes     : à partir des CSV sources (DSN, PAIE_AUDIT.csv), année complète —
    pour vérifier que les pipelines de reconstruction ne perdent aucun salarié en route ;
  - T1 (isopérimètre)  : PAIE restreinte au 1er trimestre 2026, pour comparer à la même
    fenêtre temporelle que le DSN (qui ne couvre que Q1 2026).

Clé de rapprochement : le matricule, normalisé (zéros de tête retirés) car la DSN le
stocke sur 6 chiffres ("005005") et la PAIE non ("5005"). Le NIR n'est pas utilisable
côté PAIE (valeur factice "0" sur toutes les lignes de PAIE_AUDIT.csv).
"""
import os
import shutil
import tempfile
import openpyxl
import pandas as pd
from src.normalisation import _norm_soc
from config import OUTPUT_DIR


def nom_fichier_dsn(societe):
    return f"Reconstruit_{_norm_soc(societe).replace(' ', '-')}_CORRIGE.xlsx"


def nom_fichier_paie(code):
    return f"Reconstruit_PAIE_{code}.xlsx"


def _normaliser_matricule(v):
    """Retire les zéros de tête (format DSN "005005" -> "5005", format PAIE déjà nu)."""
    v = str(v).strip()
    return v.lstrip("0") or "0"


# ===============================================================
# Chargement des populations — niveau "fichiers reconstruits"
# ===============================================================
def _parser_classeur(wb, col_matricule, col_nom, col_prenom):
    """Lit {matricule normalisé: (nom, prénom)} depuis un classeur déjà ouvert.
    Ferme le classeur avant de renvoyer le résultat (important pour le cas de la
    copie temporaire : le dossier temporaire ne peut être nettoyé tant que le
    fichier y est encore ouvert)."""
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows))
    i_mat = header.index(col_matricule)
    i_nom = header.index(col_nom) if col_nom and col_nom in header else None
    i_prenom = header.index(col_prenom) if col_prenom and col_prenom in header else None

    salaries = {}
    for r in rows:
        if len(r) <= i_mat or not r[i_mat]:
            continue
        mat = _normaliser_matricule(r[i_mat])
        nom = r[i_nom] if i_nom is not None and len(r) > i_nom else ""
        prenom = r[i_prenom] if i_prenom is not None and len(r) > i_prenom else ""
        salaries.setdefault(mat, (nom, prenom))
    wb.close()
    return salaries


def lire_salaries(path, col_matricule, col_nom=None, col_prenom=None):
    """Lit {matricule normalisé: (nom, prénom)} depuis un classeur reconstruit.
    Contourne un éventuel verrou de synchronisation en retentant sur une copie
    temporaire si la lecture directe échoue."""
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
        return _parser_classeur(wb, col_matricule, col_nom, col_prenom)
    except PermissionError:
        with tempfile.TemporaryDirectory() as tmp:
            copie = os.path.join(tmp, os.path.basename(path))
            shutil.copy2(path, copie)
            wb = openpyxl.load_workbook(copie, read_only=True)
            return _parser_classeur(wb, col_matricule, col_nom, col_prenom)


def charger_dsn_reconstruits(mapping):
    """{société DSN: {matricule: (nom, prénom)}} depuis les classeurs DSN reconstruits."""
    return {societe_dsn: lire_salaries(os.path.join(OUTPUT_DIR, nom_fichier_dsn(societe_dsn)),
                                       "Matricule", "Nom", "Prénom")
            for societe_dsn in mapping.values()}


def charger_paie_reconstruits(mapping):
    """{code PAIE: {matricule: (nom, prénom)}} depuis les classeurs PAIE reconstruits."""
    return {code_paie: lire_salaries(os.path.join(OUTPUT_DIR, nom_fichier_paie(code_paie)),
                                     "Matricule", "Nom", "Prenom")
            for code_paie in mapping.keys()}


# ===============================================================
# Chargement des populations — niveau "fichiers sources bruts"
# ===============================================================
def lire_population_dsn_source(path):
    """Population brute par société directement depuis le CSV source DSN (avant toute
    reconstruction) : {société normalisée: {matricule normalisé: (nom, prénom)}} —
    tout salarié qui apparaît au moins une fois, y compris annulations/dates invalides."""
    from src.chargement import lire_source, mapper_colonnes, normaliser_et_detecter
    df = lire_source(path)
    df.columns = df.columns.str.strip()
    cols = mapper_colonnes(df)
    df = normaliser_et_detecter(df, cols)
    df = df.loc[df["_soc"] != ""]

    pop = {}
    for soc, g in df.groupby("_soc"):
        g2 = g.drop_duplicates(subset=["Matricule"])
        salaries = {}
        for mat, nom, prenom in zip(g2["Matricule"], g2["Nom"], g2["Prénom"]):
            m = str(mat).strip()
            if not m:
                continue
            salaries[_normaliser_matricule(m)] = (nom, prenom)
        pop[soc] = salaries
    return pop


def lire_population_paie_source(path, annee, sep, date_debut=None, date_fin=None):
    """Population brute par entreprise directement depuis PAIE_AUDIT.csv (année filtrée,
    avant toute jointure PPU) : {entreprise: {matricule normalisé: (nom, prénom)}}.
    Si date_debut/date_fin sont fournies (colonne "periode"), restreint en plus à cette
    fenêtre — sert à comparer à isopérimètre avec le DSN, qui ne couvre que Q1 2026."""
    from src.chargement_paie import lire_paie, decouper_par_entreprise
    df = lire_paie(path, annee=annee, sep=sep)
    if date_debut is not None and date_fin is not None:
        dates = pd.to_datetime(df["periode"], errors="coerce")
        df = df.loc[(dates >= date_debut) & (dates <= date_fin)]

    pop = {}
    for ent, g in decouper_par_entreprise(df).items():
        g2 = g.drop_duplicates(subset=["matricule"])
        salaries = {}
        for mat, nom, prenom in zip(g2["matricule"], g2["nom"], g2["prenom"]):
            m = str(mat).strip()
            if not m:
                continue
            salaries[_normaliser_matricule(m)] = (nom, prenom)
        pop[ent] = salaries
    return pop


# ===============================================================
# Comparaison (générique, commune à tous les niveaux)
# ===============================================================
def _comparer_deux_populations(dsn_pop, paie_pop, societe_dsn, code_paie):
    """Compare deux {matricule: (nom, prénom)} pour une paire de sociétés.
    Renvoie (ligne_synthese, lignes_detail)."""
    manquants = sorted(set(dsn_pop) - set(paie_pop))
    n_communs = len(dsn_pop) - len(manquants)

    ligne_synthese = {
        "Société DSN": societe_dsn,
        "Société PAIE": code_paie,
        "Salariés PAIE (population payée)": len(paie_pop),
        "Salariés DSN (avec arrêt)": len(dsn_pop),
        "Communs": n_communs,
        "DSN sans PAIE (à vérifier)": len(manquants),
        "PAIE sans arrêt DSN": len(paie_pop) - n_communs,
        "Taux de correspondance DSN->PAIE": (n_communs / len(dsn_pop)) if dsn_pop else None,
    }

    lignes_detail = []
    for mat in sorted(set(dsn_pop) | set(paie_pop)):
        dans_dsn = mat in dsn_pop
        dans_paie = mat in paie_pop
        nom, prenom = dsn_pop.get(mat) or paie_pop.get(mat) or ("", "")
        if dans_dsn and dans_paie:
            statut = "OK - présent des deux côtés"
        elif dans_dsn and not dans_paie:
            statut = "🔴 DSN sans PAIE (à vérifier)"
        else:
            statut = "PAIE sans arrêt DSN"
        lignes_detail.append({
            "Société DSN": societe_dsn, "Société PAIE": code_paie,
            "Matricule": mat, "Nom": nom, "Prénom": prenom,
            "Présent DSN": "Oui" if dans_dsn else "Non",
            "Présent PAIE": "Oui" if dans_paie else "Non",
            "Statut": statut,
        })
    return ligne_synthese, lignes_detail


def comparer(mapping, dsn_pop_par_societe, paie_pop_par_code):
    """Compare génériquement deux populations déjà chargées, quel que soit leur niveau
    (reconstruit, source, restreint à une période...). Renvoie (synthese, detail)."""
    synthese, detail = [], []
    for code_paie, societe_dsn in mapping.items():
        dsn_pop = dsn_pop_par_societe.get(societe_dsn, {})
        paie_pop = paie_pop_par_code.get(code_paie, {})
        s, d = _comparer_deux_populations(dsn_pop, paie_pop, societe_dsn, code_paie)
        synthese.append(s)
        detail.extend(d)
    return synthese, detail
