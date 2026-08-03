# -*- coding: utf-8 -*-
"""Comparaison des jours d'absence PAIE <-> DSN, par société / salarié / mois / motif.

Un seul fichier pour toute la fonctionnalité (identification des rubriques PAIE,
extraction DSN, jointure, export) : à consulter de haut en bas, section par section.

  1. IDENTIFICATION DES RUBRIQUES PAIE — quelles rubriques comptent comme un motif
     d'absence, et lecture de la colonne "Base" (qui n'est un nombre de jours que pour
     certaines rubriques, cf. config.MOTIFS_PAIE_JOURS). Sert à construire_reconstruit()
     (src/chargement_paie.py) — pas à ce module directement (cf. point 2).
  2. EXTRACTION PAIE ET DSN — lecture des classeurs déjà reconstruits (output/xlsx/,
     produits par main_paie.py et main.py), PAS du CSV source : les reconstruits sont
     le livrable final, censés avoir déjà corrigé les anomalies identifiées — toute
     analyse en aval doit en partir, pas recalculer indépendamment depuis la donnée
     brute (cf. échanges du 2026-07-31). Reconstruit_PAIE_<code>.xlsx porte déjà
     "Motif"/"Jours d'absence (retenu)" (calculés par construire_reconstruit avec la
     même logique que la section 1) ; Reconstruit_<société>_CORRIGE.xlsx porte déjà
     "Jours absence (mois)" par motif.
  3. COMPARAISON — jointure PAIE <-> DSN sur (matricule normalisé, mois, motif).
  4. EXPORT — classeur Excel avec les colonnes d'identité demandées, une clé pivot par
     salarié, et un filtre Excel cliquable sur chaque colonne (dont l'écart).

Les deux reconstruits (PAIE et DSN) sont restreints à la même période (cf.
config.DATE_DEBUT_PERIODE/DATE_FIN_PERIODE, actuellement l'année 2026 complète) —
même périmètre des deux côtés, donc plus besoin de filtrer les mois nous-mêmes ici.
"""
import os
import re
import shutil
import tempfile
import unicodedata
import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from config import MOTIFS_PAIE_JOURS, OUTPUT_DIR, RAPPORT_DIR
from src.normalisation import _norm_soc, to_num

def nom_fichier_comparatif_absences(code, date_sortie):
    """Un classeur par société (cf. demande du 2026-08-03, cohérent avec les autres
    livrables PAIE déjà séparés par société) : Comparatif_jours_absence_PAIE_DSN_<code>_
    <date_sortie>.xlsx. `date_sortie` : date du jour de génération au format JJ-MM-AAAA."""
    return f"Comparatif_jours_absence_PAIE_DSN_{code}_{date_sortie}.xlsx"


# ===============================================================
# 0. UTILITAIRES PARTAGÉS
# ===============================================================
def normaliser_matricule(v):
    """Retire les zéros de tête (format DSN "005005" -> "5005", format PAIE déjà nu)."""
    v = str(v).strip()
    return v.lstrip("0") or "0"


def nom_fichier_dsn(societe):
    return f"Reconstruit_{_norm_soc(societe).replace(' ', '-')}_CORRIGE.xlsx"


def nom_fichier_paie(code):
    return f"Reconstruit_PAIE_{code}.xlsx"


def _lire_classeur_verrouille(path, parseur):
    """Ouvre un classeur en lecture seule et applique `parseur(wb)`. Contourne un
    éventuel verrou de synchronisation (OneDrive...) en retentant sur une copie
    temporaire — la lecture doit se terminer (et le classeur être fermé) avant la
    sortie du bloc `with`, sinon le nettoyage du dossier temporaire échoue lui-même."""
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
        return parseur(wb)
    except PermissionError:
        with tempfile.TemporaryDirectory() as tmp:
            copie = os.path.join(tmp, os.path.basename(path))
            shutil.copy2(path, copie)
            wb = openpyxl.load_workbook(copie, read_only=True)
            return parseur(wb)


# ===============================================================
# 1. IDENTIFICATION DES RUBRIQUES PAIE ET JOURS D'ABSENCE
# ===============================================================
_PONCTUATION = re.compile(r"[-/.:',]")


def _normaliser_libelle(libelle):
    """Minuscules, sans accents, ponctuation remplacée par des espaces, espaces
    multiples réduits à un seul — tolère les variantes de transcription du référentiel
    métier (ex. "CP-Abs..." et "CP - Abs..." ou "Maternite"/"Maternité" sont identiques)."""
    s = unicodedata.normalize("NFD", str(libelle))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = _PONCTUATION.sub(" ", s.lower())
    return " ".join(s.split())


def _construire_lookup_motifs(mapping=MOTIFS_PAIE_JOURS):
    """{(rubrique, libellé normalisé): {motifs}} à partir d'un mapping
    motif -> [(libellé, code), ...] (par défaut config.MOTIFS_PAIE_JOURS)."""
    lookup = {}
    for motif, entrees in mapping.items():
        for libelle, code in entrees:
            cle = (str(code).strip(), _normaliser_libelle(libelle))
            lookup.setdefault(cle, set()).add(motif)
    return lookup


def identifier_motifs(df):
    """Ajoute une colonne "_motifs" (set des motifs reconnus, vide si aucune
    correspondance) au DataFrame PAIE (colonnes source : rubrique, libelle).
    Fonction publique : réutilisée par src/chargement_paie.py (construire_reconstruit)
    pour que Reconstruit_PAIE_<code>.xlsx porte lui-même le motif/jours d'absence
    retenu, plutôt que de laisser cette info exister seulement dans ce module."""
    lookup = _construire_lookup_motifs()
    libelles_norm = df["libelle"].apply(_normaliser_libelle)
    cles = list(zip(df["rubrique"].astype(str).str.strip(), libelles_norm))
    df = df.copy()
    df["_motifs"] = [lookup.get(c, set()) for c in cles]
    return df


_COLONNES_RECONSTRUIT_PAIE = ("Matricule", "Nom", "Prenom", "Entreprise", "Etablissement",
                              "Siren", "Nic", "Siret", "Periode", "Motif",
                              "Jours d'absence (retenu)")


def _parser_reconstruit_paie(wb):
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows))
    idx = {c: header.index(c) for c in _COLONNES_RECONSTRUIT_PAIE}
    lignes = []
    for r in rows:
        motif, jours = r[idx["Motif"]], r[idx["Jours d'absence (retenu)"]]
        if not motif or jours in (None, ""):
            continue
        lignes.append({
            "Matricule": normaliser_matricule(r[idx["Matricule"]]),
            "Nom": r[idx["Nom"]], "Prenom": r[idx["Prenom"]],
            "Entreprise": r[idx["Entreprise"]], "Etablissement": r[idx["Etablissement"]] or "",
            "Siren": r[idx["Siren"]] or "", "Nic": r[idx["Nic"]] or "", "Siret": r[idx["Siret"]] or "",
            "Mois": str(r[idx["Periode"]])[:7], "Motif": motif, "Jours": float(jours),
        })
    wb.close()
    return pd.DataFrame(lignes)


def charger_jours_paie(mapping):
    """{code entreprise: DataFrame(Matricule, Nom, Prenom, Etablissement, Siren, Nic,
    Siret, Mois, Motif, Jours)} lu directement dans Reconstruit_PAIE_<code>.xlsx —
    "Motif"/"Jours d'absence (retenu)" y sont déjà calculés par construire_reconstruit()
    (src/chargement_paie.py, même logique qu'ici, cf. identifier_motifs). Le reconstruit
    est le livrable final : on part de lui, pas du CSV source (cf. docstring module).

    "Motif" peut contenir plusieurs valeurs séparées par ", " (rubrique comptant pour
    plusieurs motifs à la fois, cf. config.MOTIFS_PAIE_JOURS) : la ligne est alors
    dupliquée (explode), une fois par motif."""
    resultat = {}
    for code_paie in mapping.keys():
        path = os.path.join(OUTPUT_DIR, nom_fichier_paie(code_paie))
        detail = _lire_classeur_verrouille(path, _parser_reconstruit_paie)
        if detail.empty:
            resultat[code_paie] = detail
            continue
        detail = detail.assign(Motif=detail["Motif"].str.split(", ")).explode("Motif")
        resultat[code_paie] = (detail.groupby(
            ["Matricule", "Nom", "Prenom", "Entreprise", "Etablissement", "Siren", "Nic", "Siret",
             "Mois", "Motif"], as_index=False)["Jours"].sum())
    return resultat


COLONNES_ANOMALIES = ["Entreprise", "Matricule", "Nom", "Prenom", "Mois", "Motif",
                      "Rubrique", "Libelle", "Base", "Anomalie"]


def detecter_anomalies_jours(chemin_paie_source, annee, sep):
    """Signale les lignes PAIE_AUDIT (sur les rubriques "jours" de MOTIFS_PAIE_JOURS)
    où "Base" est physiquement incohérente pour un mois calendaire : plus de 31 jours,
    ou négative sur une rubrique qui n'est pas une annulation (le libellé contient
    "annul"). Confirmé sur la donnée elle-même (pas un artefact d'agrégation) : voir
    échanges de validation — chaque cas correspond à une seule ligne brute du CSV
    source, sans information de sous-période/rétroactivité exploitable pour l'expliquer
    (date_sous_periode/num_sous_periode/date_retro sont à blanc sur ces lignes)."""
    from src.chargement_paie import lire_paie
    df = lire_paie(chemin_paie_source, annee=annee, sep=sep)
    df = identifier_motifs(df)
    df = df.loc[df["_motifs"].apply(len) > 0].copy()
    if df.empty:
        return pd.DataFrame(columns=COLONNES_ANOMALIES)

    df["Base_num"] = to_num(df["base"])
    est_annulation = df["libelle"].apply(_normaliser_libelle).str.contains("annul")
    anormal = (df["Base_num"] > 31) | ((df["Base_num"] < 0) & ~est_annulation)
    df = df.loc[anormal].copy()
    if df.empty:
        return pd.DataFrame(columns=COLONNES_ANOMALIES)

    df["_motifs"] = df["_motifs"].apply(list)
    df = df.explode("_motifs")
    df["Matricule"] = df["matricule"].apply(normaliser_matricule)
    df["Mois"] = df["periode"].str[:7]
    df["Anomalie"] = df["Base_num"].apply(
        lambda b: "Base > 31 j sur un mois" if b > 31 else "Base négative (hors annulation)")

    sortie = pd.DataFrame({
        "Entreprise": df["entreprise"], "Matricule": df["Matricule"], "Nom": df["nom"],
        "Prenom": df["prenom"], "Mois": df["Mois"], "Motif": df["_motifs"],
        "Rubrique": df["rubrique"], "Libelle": df["libelle"], "Base": df["Base_num"],
        "Anomalie": df["Anomalie"],
    })
    return sortie.sort_values(["Entreprise", "Matricule", "Mois"]).reset_index(drop=True)


# ===============================================================
# 2. EXTRACTION DSN (classeurs déjà reconstruits)
# ===============================================================
_COLONNES_DSN = ("Matricule", "Nom", "Prénom", "Nom usage", "NIR", "Motif arrêt",
                "Mois DSN", "Jours absence (mois)", "Type de ligne")
_LIGNE_DONNEES = "📍 Données DSN"


def canoniser_motif_dsn(motif):
    """Ramène un motif DSN (texte libre) à l'un des paniers alignés sur le référentiel
    PAIE (config.MOTIFS_PAIE_JOURS) : regroupe accident de travail + accident de trajet
    + maladie professionnelle (non distingués côté paie). Les motifs sans équivalent
    PAIE identifié à ce jour (temps partiel thérapeutique, adoption) gardent leur
    propre étiquette plutôt que d'être rattachés arbitrairement à un autre."""
    m = str(motif).strip().lower()
    if "accident de travail" in m or "accident du travail" in m:
        return "accident de travail"
    if "accident de trajet" in m or "accident du trajet" in m:
        return "accident de travail"
    if "maladie professionnelle" in m:
        return "accident de travail"
    if "maternité" in m:
        return "maternité"
    if "paternité" in m:
        return "paternité"
    if "adoption" in m:
        return "adoption (pas de rubrique PAIE identifiée)"
    if "temps partiel" in m:
        return "temps partiel thérapeutique (pas de rubrique PAIE identifiée)"
    if "maladie" in m:
        return "maladie"
    return "autre / non classé"


def _parser_dsn(wb):
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows))
    idx = {c: header.index(c) for c in _COLONNES_DSN}
    lignes = []
    for r in rows:
        mat, mois, jours = r[idx["Matricule"]], r[idx["Mois DSN"]], r[idx["Jours absence (mois)"]]
        if not mat or not mois or jours in (None, "") or r[idx["Type de ligne"]] != _LIGNE_DONNEES:
            continue
        lignes.append({
            "Matricule": normaliser_matricule(mat), "Nom": r[idx["Nom"]], "Prenom": r[idx["Prénom"]],
            "Nom usage": r[idx["Nom usage"]] or "", "NIR": r[idx["NIR"]] or "",
            "Mois": str(mois), "Motif": canoniser_motif_dsn(r[idx["Motif arrêt"]]),
            "Jours": float(jours),
        })
    wb.close()
    return pd.DataFrame(lignes)


def charger_jours_dsn(mapping):
    """{société DSN: DataFrame(Matricule, Nom, Prenom, Nom usage, NIR, Mois, Motif,
    Jours)} depuis les classeurs DSN déjà reconstruits. Seules les lignes "📍 Données
    DSN" sont sommées : les lignes "❌ Annulation DSN" sont un historique de corrections
    déjà intégrées dans le résultat positif survivant (cf. pipeline.py, survivance par
    (emp, motif, DJT)) — les ressommer reviendrait à soustraire deux fois les mêmes
    corrections."""
    resultat = {}
    for societe_dsn in mapping.values():
        path = os.path.join(OUTPUT_DIR, nom_fichier_dsn(societe_dsn))
        detail = _lire_classeur_verrouille(path, _parser_dsn)
        if detail.empty:
            resultat[societe_dsn] = detail
            continue
        resultat[societe_dsn] = (detail.groupby(
            ["Matricule", "Nom", "Prenom", "Nom usage", "NIR", "Mois", "Motif"],
            as_index=False)["Jours"].sum())
    return resultat


_COL_SIREN_DSN = "S21.G00.06.001 Entreprise.Siren"
_COL_NIC_DSN = "S21.G00.11.001 Etablissement.Nic"


def lire_identite_dsn_source(chemin_dsn_source):
    """{matricule normalisé: (siren, nic, siret)} depuis le CSV source DSN (pas le
    classeur reconstruit, qui ne porte pas ces colonnes). Le SIREN est unique par
    société et le NIC constant par salarié (vérifié empiriquement sur 2026 : aucun
    matricule n'a plus d'un NIC) ; le SIRET est simplement SIREN+NIC concaténés (14
    chiffres). Bien plus fiable que la PAIE, où ces 3 champs sont à 0 pour toutes les
    lignes de PAIE_AUDIT.csv."""
    from src.chargement import lire_source, mapper_colonnes, normaliser_et_detecter
    df = lire_source(chemin_dsn_source)
    df.columns = df.columns.str.strip()
    cols = mapper_colonnes(df)
    df = normaliser_et_detecter(df, cols)
    df = df.loc[(df["_soc"] != "") & df[_COL_SIREN_DSN].notna() & df[_COL_NIC_DSN].notna()]
    df = df.drop_duplicates(subset=["Matricule"])

    identite = {}
    for mat, siren, nic in zip(df["Matricule"], df[_COL_SIREN_DSN], df[_COL_NIC_DSN]):
        siren, nic = str(siren).strip(), str(nic).strip()
        identite[normaliser_matricule(mat)] = (siren, nic, f"{siren}{nic}")
    return identite


# ===============================================================
# 3. COMPARAISON
# ===============================================================
COLONNES_SORTIE = [
    "Clé salarié (pivot)", "Nom", "Prénom", "Nom d'usage", "NIR", "Matricule",
    "Entreprise", "Etablissement", "Siren", "Nic", "Siret", "Motif",
    "Nombre de jours absence PAIE", "Nombre de jours absence DSN",
    "Mois absence (période)", "Écart constaté",
]


def comparer_jours(mapping, jours_dsn_par_societe, jours_paie_par_code, identite_dsn=None):
    """Fusionne DSN et PAIE par société sur (Matricule, Mois, Motif). Renvoie un
    DataFrame avec les colonnes COLONNES_SORTIE — une ligne par (société, salarié,
    mois, motif). La "Clé salarié (pivot)" (Entreprise + Matricule) permet de vérifier
    d'un coup d'œil qu'on compare bien le même salarié des deux côtés.

    `identite_dsn` (cf. lire_identite_dsn_source) : {matricule: (siren, nic, siret)}
    — si fourni, écrase le Siren/Nic/Siret de la PAIE (toujours à 0 dans PAIE_AUDIT.csv)
    par la valeur réelle issue de la DSN, quand le salarié y est identifié."""
    identite_dsn = identite_dsn or {}
    lignes = []
    cols_vides = ["Matricule", "Nom", "Prenom", "Mois", "Motif", "Jours"]
    for code_paie, societe_dsn in mapping.items():
        dsn = jours_dsn_par_societe.get(societe_dsn)
        if dsn is None or dsn.empty:
            dsn = pd.DataFrame(columns=cols_vides + ["Nom usage", "NIR"])
        paie = jours_paie_par_code.get(code_paie)
        if paie is None or paie.empty:
            paie = pd.DataFrame(columns=cols_vides + ["Etablissement", "Siren", "Nic", "Siret"])

        fusion = pd.merge(dsn, paie, on=["Matricule", "Mois", "Motif"], how="outer",
                          suffixes=("_DSN", "_PAIE"))
        fusion["Nom"] = fusion["Nom_DSN"].fillna(fusion["Nom_PAIE"])
        fusion["Prénom"] = fusion["Prenom_DSN"].fillna(fusion["Prenom_PAIE"])
        fusion["Jours_DSN"] = fusion["Jours_DSN"].fillna(0.0)
        fusion["Jours_PAIE"] = fusion["Jours_PAIE"].fillna(0.0)

        fusion["Clé salarié (pivot)"] = f"{code_paie}-" + fusion["Matricule"].astype(str)
        fusion["Entreprise"] = code_paie
        for col in ("Nom usage", "NIR", "Etablissement", "Siren", "Nic", "Siret"):
            if col not in fusion.columns:
                fusion[col] = ""
            else:
                fusion[col] = fusion[col].fillna("")

        # Nom usage/NIR/Etablissement ne viennent que d'un des deux côtés (DSN ou PAIE) :
        # sur un mois où ce côté n'a pas de ligne pour ce salarié, la valeur reste vide,
        # alors qu'un autre mois du même salarié peut l'avoir. Sans cette propagation,
        # comparer_jours_periode() (qui regroupe notamment sur ces colonnes) fractionne à
        # tort un même salarié en plusieurs lignes selon les mois où chaque source a parlé
        # (vu concrètement sur TEYSSEYRE Justine/AURA : Etablissement vide en janvier,
        # "2" en février, scindant un total par ailleurs strictement identique 21=21).
        for col in ("Nom usage", "NIR", "Etablissement"):
            premiere_valeur = fusion.groupby("Matricule")[col].transform(
                lambda s: next((v for v in s if v not in (None, "")), ""))
            fusion[col] = premiere_valeur

        if identite_dsn:
            reel = fusion["Matricule"].map(identite_dsn)
            a_identite = reel.notna()
            fusion.loc[a_identite, "Siren"] = reel[a_identite].apply(lambda t: t[0])
            fusion.loc[a_identite, "Nic"] = reel[a_identite].apply(lambda t: t[1])
            fusion.loc[a_identite, "Siret"] = reel[a_identite].apply(lambda t: t[2])

        lignes.append(pd.DataFrame({
            "Clé salarié (pivot)": fusion["Clé salarié (pivot)"],
            "Nom": fusion["Nom"], "Prénom": fusion["Prénom"],
            "Nom d'usage": fusion["Nom usage"], "NIR": fusion["NIR"],
            "Matricule": fusion["Matricule"], "Entreprise": fusion["Entreprise"],
            "Etablissement": fusion["Etablissement"], "Siren": fusion["Siren"],
            "Nic": fusion["Nic"], "Siret": fusion["Siret"], "Motif": fusion["Motif"],
            "Nombre de jours absence PAIE": fusion["Jours_PAIE"],
            "Nombre de jours absence DSN": fusion["Jours_DSN"],
            "Mois absence (période)": fusion["Mois"],
            "Écart constaté": fusion["Jours_DSN"] - fusion["Jours_PAIE"],
        }))

    return pd.concat(lignes, ignore_index=True) if lignes else pd.DataFrame(columns=COLONNES_SORTIE)


def comparer_jours_periode(detail_mensuel, mois_retenus, libelle_periode="T1"):
    """Agrège le détail mensuel sur une période (ex. le T1 : ["2026-01", "2026-02",
    "2026-03"]) par (Clé salarié, Motif) — même colonnes que le détail mensuel, avec
    "Mois absence (période)" remplacé par le libellé de la période agrégée."""
    sous = detail_mensuel.loc[detail_mensuel["Mois absence (période)"].isin(mois_retenus)]
    cols_identite = ["Clé salarié (pivot)", "Nom", "Prénom", "Nom d'usage", "NIR", "Matricule",
                     "Entreprise", "Etablissement", "Siren", "Nic", "Siret", "Motif"]
    agg = (sous.groupby(cols_identite, as_index=False)
              [["Nombre de jours absence PAIE", "Nombre de jours absence DSN"]].sum())
    agg["Mois absence (période)"] = libelle_periode
    agg["Écart constaté"] = agg["Nombre de jours absence DSN"] - agg["Nombre de jours absence PAIE"]
    return agg[COLONNES_SORTIE]


COLONNES_TRACABILITE = COLONNES_SORTIE + ["Total jours DSN (mois)", "Total jours PAIE (mois)",
                                          "Écart total (mois)"]


def construire_synthese_ok(mapping, detail_mensuel):
    """Synthèse des jours d'arrêt DSN vs PAIE par société et par mois, calculée
    DIRECTEMENT à partir du détail mensuel (onglet "1 - Détail mensuel", le DataFrame
    produit par comparer_jours) — pas de population reconstruite indépendamment (ex.
    tous les salariés "communs" aux deux sources sur toute la période, complétée par des
    0 = 0 pour les mois sans absence). La population de référence pour un (société,
    mois) donné est exactement celle qui a au moins une ligne dans le détail mensuel
    pour ce (société, mois) — un salarié sans aucune ligne ce mois-là n'est pas compté.
    Les deux onglets restent ainsi strictement cohérents (cf. échange du 2026-07-31:
    "l'onglet synthèse devrait partir du détail mensuel et refléter la réalité").

    Un salarié est en écart un mois donné quand son total mensuel (DSN vs PAIE, tous
    motifs confondus ce mois-là) ne correspond pas exactement — même critère que la
    colonne "Écart constaté", simplement agrégé au niveau salarié/mois plutôt que
    salarié/mois/motif (motifs hors périmètre ici, cf. demande du 2026-07-31). Ce critère
    sert à construire la traçabilité (ci-dessous) ; la synthèse elle-même totalise les
    jours, elle ne compte plus les salariés individuellement en OK/Pas OK (cf. demande du
    2026-08-03 : partir des vrais volumes de jours plutôt que d'un comptage par salarié).

    Renvoie (synthese, tracabilite) :
      - synthese : une ligne par (société, mois) — Mois, Salariés en communs, Nombre de
        jours d'arrêt en DSN, Nombre de jours d'arrêt en PAIE, Écart constaté (somme des
        écarts individuels du mois) — un tableau par société une fois exporté ;
      - tracabilite : le détail mensuel (grain motif), restreint aux (salarié, mois) où
        le total ne correspond pas, avec en plus le total DSN/PAIE/écart du mois pour ce
        salarié — pour permettre la vérification / correction / arbitrage au cas par cas.
    """
    colonnes_synthese = ["Entreprise", "Société DSN", "Mois", "Salariés en communs",
                         "Nombre de jours d'arrêt en DSN", "Nombre de jours d'arrêt en PAIE",
                         "Écart constaté"]
    if detail_mensuel.empty:
        return (pd.DataFrame(columns=colonnes_synthese), pd.DataFrame(columns=COLONNES_TRACABILITE))

    agg = (detail_mensuel.groupby(["Entreprise", "Matricule", "Mois absence (période)"], as_index=False)
                        [["Nombre de jours absence PAIE", "Nombre de jours absence DSN"]].sum())
    agg["Écart total (mois)"] = agg["Nombre de jours absence DSN"] - agg["Nombre de jours absence PAIE"]
    agg["OK"] = agg["Écart total (mois)"].abs() < 1e-6

    lignes_synthese = []
    for (entreprise, mois), groupe in agg.groupby(["Entreprise", "Mois absence (période)"]):
        lignes_synthese.append({
            "Entreprise": entreprise, "Société DSN": mapping.get(entreprise, ""), "Mois": mois,
            "Salariés en communs": len(groupe),
            "Nombre de jours d'arrêt en DSN": groupe["Nombre de jours absence DSN"].sum(),
            "Nombre de jours d'arrêt en PAIE": groupe["Nombre de jours absence PAIE"].sum(),
            "Écart constaté": groupe["Écart total (mois)"].sum(),
        })
    synthese = pd.DataFrame(lignes_synthese, columns=colonnes_synthese).sort_values(
        ["Entreprise", "Mois"]).reset_index(drop=True)

    pas_ok = agg.loc[~agg["OK"], ["Entreprise", "Matricule", "Mois absence (période)",
                                  "Nombre de jours absence DSN", "Nombre de jours absence PAIE",
                                  "Écart total (mois)"]].rename(columns={
        "Nombre de jours absence DSN": "Total jours DSN (mois)",
        "Nombre de jours absence PAIE": "Total jours PAIE (mois)",
    })
    if not pas_ok.empty:
        tracabilite = pd.merge(detail_mensuel, pas_ok,
                               on=["Entreprise", "Matricule", "Mois absence (période)"], how="inner")
        tracabilite = tracabilite[COLONNES_TRACABILITE].sort_values(
            ["Entreprise", "Matricule", "Mois absence (période)", "Motif"]).reset_index(drop=True)
    else:
        tracabilite = pd.DataFrame(columns=COLONNES_TRACABILITE)

    return synthese, tracabilite


# ===============================================================
# 4. EXPORT EXCEL
# ===============================================================
NAVY = "1F3864"
REDF = "F4D6D2"
GREENF = "DDEBE2"
ORANGEF = "FBE7CE"


def _styler_entete_et_filtre(ws, df):
    for cell in ws[1]:
        cell.font = Font(bold=True, size=10, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
    ws.freeze_panes = "A2"
    # Filtre Excel cliquable sur chaque colonne : l'utilisateur ouvre le menu déroulant
    # de l'en-tête pour isoler ce qui l'intéresse (ex. "Écart constaté" différent de 0).
    ws.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}"


def _ecrire_onglet(writer, df, nom_onglet):
    df = df if not df.empty else pd.DataFrame(columns=COLONNES_SORTIE)
    df.to_excel(writer, index=False, sheet_name=nom_onglet)
    ws = writer.sheets[nom_onglet]
    _styler_entete_et_filtre(ws, df)

    j_ecart = list(df.columns).index("Écart constaté") + 1
    for i, ecart in enumerate(df["Écart constaté"], start=2):
        ws.cell(i, j_ecart).fill = PatternFill("solid", fgColor=(GREENF if ecart == 0 else REDF))


def _ecrire_onglet_anomalies(writer, df, nom_onglet):
    df = df if not df.empty else pd.DataFrame(columns=COLONNES_ANOMALIES)
    df.to_excel(writer, index=False, sheet_name=nom_onglet)
    ws = writer.sheets[nom_onglet]
    _styler_entete_et_filtre(ws, df)
    for i in range(2, len(df) + 2):
        for j in range(1, len(df.columns) + 1):
            ws.cell(i, j).fill = PatternFill("solid", fgColor=ORANGEF)


def _ecrire_onglet_synthese(writer, synthese, nom_onglet):
    """Un tableau par société (empilés dans le même onglet, séparés par une ligne
    vide) : Mois / Salariés en communs / Nombre de jours d'arrêt en DSN / Nombre de
    jours d'arrêt en PAIE / Écart constaté. "Salariés en communs" = ceux qui
    apparaissent dans le détail mensuel ce mois-là (cf. construire_synthese_ok). Vert
    si l'écart du mois est nul, rouge sinon."""
    wb = writer.book
    ws = wb.create_sheet(nom_onglet)
    entetes = ["Mois", "Salariés en communs", "Nombre de jours d'arrêt en DSN",
              "Nombre de jours d'arrêt en PAIE", "Écart constaté"]
    r = 1
    if synthese.empty:
        for j, h in enumerate(entetes, start=1):
            ws.cell(r, j, h)
    else:
        for entreprise, groupe in synthese.groupby("Entreprise", sort=False):
            societe_dsn = groupe["Société DSN"].iloc[0]
            titre = ws.cell(r, 1, f"{entreprise} ({societe_dsn})")
            titre.font = Font(bold=True, size=11, color="FFFFFF")
            for j in range(1, len(entetes) + 1):
                ws.cell(r, j).fill = PatternFill("solid", fgColor=NAVY)
            r += 1
            for j, h in enumerate(entetes, start=1):
                c = ws.cell(r, j, h)
                c.font = Font(bold=True, size=10)
                c.fill = PatternFill("solid", fgColor="D9D9D9")
            r += 1
            for _, ligne in groupe.iterrows():
                ws.cell(r, 1, ligne["Mois"])
                ws.cell(r, 2, int(ligne["Salariés en communs"]))
                ws.cell(r, 3, float(ligne["Nombre de jours d'arrêt en DSN"]))
                ws.cell(r, 4, float(ligne["Nombre de jours d'arrêt en PAIE"]))
                ws.cell(r, 5, float(ligne["Écart constaté"]))
                fond = GREENF if ligne["Écart constaté"] == 0 else REDF
                for j in range(1, len(entetes) + 1):
                    ws.cell(r, j).fill = PatternFill("solid", fgColor=fond)
                r += 1
            r += 1  # ligne vide entre deux sociétés
    for col, largeur in zip("ABCDE", (14, 18, 24, 24, 16)):
        ws.column_dimensions[col].width = largeur


def _ecrire_onglet_tracabilite(writer, tracabilite, nom_onglet):
    df = tracabilite if not tracabilite.empty else pd.DataFrame(columns=COLONNES_TRACABILITE)
    df.to_excel(writer, index=False, sheet_name=nom_onglet)
    ws = writer.sheets[nom_onglet]
    _styler_entete_et_filtre(ws, df)
    j_ecart = list(df.columns).index("Écart total (mois)") + 1
    for i in range(2, len(df) + 2):
        ws.cell(i, j_ecart).fill = PatternFill("solid", fgColor=REDF)


def ecrire_comparatif_absences(mapping, detail_mensuel, detail_periode, anomalies, synthese_ok,
                               tracabilite_ecarts, date_sortie, libelle_periode="T1"):
    """Écrit UN CLASSEUR PAR SOCIÉTÉ (cf. demande du 2026-08-03) : chaque société de
    `mapping` reçoit son propre Comparatif_jours_absence_PAIE_DSN_<code>_<date_sortie>.xlsx,
    avec les 5 mêmes onglets qu'auparavant mais filtrés sur cette société (toutes les
    DataFrames en entrée portent une colonne "Entreprise" = code PAIE). Renvoie la liste
    des chemins de fichiers écrits."""
    fichiers = []
    for code_paie in mapping.keys():
        out = os.path.join(RAPPORT_DIR, nom_fichier_comparatif_absences(code_paie, date_sortie))
        sous_detail = detail_mensuel.loc[detail_mensuel["Entreprise"] == code_paie]
        sous_periode = detail_periode.loc[detail_periode["Entreprise"] == code_paie]
        sous_anomalies = anomalies.loc[anomalies["Entreprise"] == code_paie]
        sous_synthese = synthese_ok.loc[synthese_ok["Entreprise"] == code_paie]
        sous_tracabilite = tracabilite_ecarts.loc[tracabilite_ecarts["Entreprise"] == code_paie]
        with pd.ExcelWriter(out, engine="openpyxl") as writer:
            _ecrire_onglet(writer, sous_detail, "1 - Détail mensuel")
            _ecrire_onglet(writer, sous_periode, f"2 - Total {libelle_periode}")
            _ecrire_onglet_anomalies(writer, sous_anomalies, "3 - Anomalies (jours)")
            _ecrire_onglet_synthese(writer, sous_synthese, "4 - Jours DSN vs PAIE")
            _ecrire_onglet_tracabilite(writer, sous_tracabilite, "5 - Traçabilité écarts")
        print(f"   ✅ écrit -> {out}")
        fichiers.append(out)
    return fichiers


# ===============================================================
# 5. FICHIER "JOURS D'ABSENCE" CÔTÉ PAIE (miroir du fichier DSN, livrable autonome)
# ===============================================================
COLONNES_ABSENCES_PAIE = ["Matricule", "Nom", "Prenom", "Entreprise", "Etablissement",
                          "Siren", "Nic", "Siret", "Motif", "Mois",
                          "Jours absence (mois)"]


def nom_fichier_absences_paie(code, date_sortie):
    """Livrable 1 (réunion du 2026-07-31, Solange) : « fichier nombre de jours d'absence
    côté paye, miroir de celui déjà réalisé côté DSN ». Nommage volontairement proche de
    celui du comparatif (Comparatif_jours_absence_PAIE_DSN.xlsx) pour que les deux se
    retrouvent facilement ensemble, et distinct de Reconstruit_PAIE_<code>.xlsx (qui est
    le détail brut par rubrique/PPU, pas le livrable "jours d'absence"). `date_sortie` :
    date du jour de génération au format JJ-MM-AAAA (cf. demande du 2026-08-03)."""
    return f"Jours_Absence_PAIE_{code}_{date_sortie}.xlsx"


def construire_absences_paie(jours_paie_par_code, identite_dsn=None):
    """{code entreprise: DataFrame(COLONNES_ABSENCES_PAIE)} — même grain et même mesure
    que charger_jours_paie() (une ligne par salarié/mois/motif), habillé pour être un
    livrable autonome : Siren/Nic/Siret repris de la DSN quand disponibles (cf.
    lire_identite_dsn_source), comme dans le comparatif."""
    identite_dsn = identite_dsn or {}
    resultat = {}
    for code, df in jours_paie_par_code.items():
        df = df.rename(columns={"Jours": "Jours absence (mois)"}).copy()
        for col in ("Etablissement", "Siren", "Nic", "Siret"):
            if col not in df.columns:
                df[col] = ""
            else:
                df[col] = df[col].fillna("")
        if identite_dsn:
            reel = df["Matricule"].map(identite_dsn)
            a_identite = reel.notna()
            df.loc[a_identite, "Siren"] = reel[a_identite].apply(lambda t: t[0])
            df.loc[a_identite, "Nic"] = reel[a_identite].apply(lambda t: t[1])
            df.loc[a_identite, "Siret"] = reel[a_identite].apply(lambda t: t[2])
        df["Entreprise"] = code
        resultat[code] = df[COLONNES_ABSENCES_PAIE].sort_values(["Matricule", "Mois", "Motif"])
    return resultat


def ecrire_absences_paie(absences_par_code, date_sortie):
    """Écrit un classeur par société dans output/xlsx/ (même dossier que les classeurs
    DSN et PAIE déjà reconstruits) : Jours_Absence_PAIE_<code>_<date_sortie>.xlsx."""
    fichiers = []
    for code, df in absences_par_code.items():
        out = os.path.join(OUTPUT_DIR, nom_fichier_absences_paie(code, date_sortie))
        with pd.ExcelWriter(out, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Jours absence")
            ws = writer.sheets["Jours absence"]
            _styler_entete_et_filtre(ws, df)
        print(f"   ✅ écrit -> {out}  ({len(df)} ligne(s))")
        fichiers.append(out)
    return fichiers
