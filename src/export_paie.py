# -*- coding: utf-8 -*-
"""Écriture des classeurs PAIE, un par société (étape 1 : structure uniquement).

Écriture en mode "write_only" (openpyxl) : les fichiers source atteignent plusieurs
centaines de milliers de lignes par société, un classeur standard (toutes les cellules
gardées en mémoire, comme le fait pandas.ExcelWriter/ws.cell()) provoque un MemoryError
à ce volume. Le mode write_only écrit chaque ligne directement dans le flux du fichier.
"""
import os
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from config import OUTPUT_DIR

NAVY = "1F3864"


def nom_fichier_paie(entreprise):
    return f"Reconstruit_PAIE_{entreprise}.xlsx"


def ecrire_xlsx_paie(entreprise, detail):
    """Écrit le classeur d'une société : un onglet, grain rubrique inchangé."""
    out = os.path.join(OUTPUT_DIR, nom_fichier_paie(entreprise))
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Reconstruit")
    ws.freeze_panes = "A2"

    header_font = Font(bold=True, size=10, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor=NAVY)
    entete = []
    for nom_col in detail.columns:
        c = WriteOnlyCell(ws, value=nom_col)
        c.font = header_font
        c.fill = header_fill
        entete.append(c)
    ws.append(entete)

    for ligne in detail.itertuples(index=False, name=None):
        ws.append(ligne)

    wb.save(out)
    print(f"   ✅ écrit -> {out}  ({len(detail)} ligne(s))")
    return out
