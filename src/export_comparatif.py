# -*- coding: utf-8 -*-
"""Écriture du classeur de comparaison des populations PAIE <-> DSN."""
import os
import pandas as pd
from openpyxl.styles import Font, PatternFill
from config import RAPPORT_DIR, COMPARATIF_XLSX

NAVY = "1F3864"
REDF = "F4D6D2"
GREENF = "DDEBE2"
ORANGEF = "FBE7CE"

COLS_DETAIL = ["Société DSN", "Société PAIE", "Matricule", "Nom", "Prénom",
              "Présent DSN", "Présent PAIE", "Statut"]


def _styler_entete(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True, size=10, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
    ws.freeze_panes = "A2"


def _ecrire_synthese(writer, synthese, nom_onglet):
    df = pd.DataFrame(synthese)
    df.to_excel(writer, index=False, sheet_name=nom_onglet)
    ws = writer.sheets[nom_onglet]
    _styler_entete(ws)
    for i, taux in enumerate(df["Taux de correspondance"], start=2):
        fill = GREENF if taux == 1 else REDF
        for j in range(1, len(df.columns) + 1):
            ws.cell(i, j).fill = PatternFill("solid", fgColor=fill)


def _ecrire_detail(writer, detail, nom_onglet):
    df = pd.DataFrame(detail) if detail else pd.DataFrame(columns=COLS_DETAIL)
    df.to_excel(writer, index=False, sheet_name=nom_onglet)
    ws = writer.sheets[nom_onglet]
    _styler_entete(ws)
    for i, statut in enumerate(df["Statut"], start=2):
        if statut.startswith("OK"):
            fill = GREENF
        elif statut.startswith("🔴"):
            fill = REDF
        else:
            fill = ORANGEF
        for j in range(1, len(df.columns) + 1):
            ws.cell(i, j).fill = PatternFill("solid", fgColor=fill)


def ecrire_comparatif(niveaux, date_sortie):
    """niveaux : liste de (titre_court, synthese, detail), un triplet par niveau de
    comparaison (ex. "reconstruits", "sources brutes", "T1 (reconstruits)", ...).
    Écrit 2 onglets par niveau (synthèse + détail), numérotés dans l'ordre fourni.
    `date_sortie` : date du jour de génération au format JJ-MM-AAAA, ajoutée à la fin
    du nom de fichier (cf. demande du 2026-08-03)."""
    base, ext = os.path.splitext(COMPARATIF_XLSX)
    out = os.path.join(RAPPORT_DIR, f"{base}_{date_sortie}{ext}")
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        n = 1
        for titre, synthese, detail in niveaux:
            _ecrire_synthese(writer, synthese, f"{n} - Comparatif ({titre})")
            n += 1
            _ecrire_detail(writer, detail, f"{n} - Détail ({titre})")
            n += 1

    print(f"   ✅ écrit -> {out}")
    return out
